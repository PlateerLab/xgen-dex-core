"""동봉본 검증 — **동봉된 인터프리터로** 돌린다: ``<python> -I bundle/verify.py``

1. 설치된 배포판에 금지 목록(서버 전용)이 없다.
2. 엔진·런타임·문서 파서가 import 된다.
3. 데몬이 ``ready`` 를 내고, 가짜 LLM 턴(Write · Bash · ParseDocument ×4)이 끝까지 돌아 종결 하나로 끝난다.

키도 네트워크도 쓰지 않는다. 실패하면 0 이 아닌 값으로 끝난다.
"""

from __future__ import annotations

import json
import os
import queue
import subprocess
import sys
import tempfile
import threading
import time
from importlib import metadata
from pathlib import Path

HERE = Path(__file__).resolve().parent
TERMINALS = ("done", "error", "cancelled")


def _normalize(name: str) -> str:
    return name.lower().replace("_", "-").replace(".", "-")


def check_forbidden() -> None:
    forbidden = {_normalize(n) for n in json.loads((HERE / "bundle.json").read_text(encoding="utf-8"))["forbidden"]}
    installed = {_normalize(d.metadata["Name"]) for d in metadata.distributions()}
    found = sorted(installed & forbidden)
    if found:
        raise SystemExit(f"forbidden packages in the bundle: {found}")
    print(f"forbidden: none of {len(forbidden)} present ({len(installed)} distributions installed)")


def check_imports() -> None:
    import xd_engine.daemon  # noqa: F401
    import xgen_agent_runtime.host.turn_executor  # noqa: F401
    import xgen_agent_runtime.llm_client.anthropic  # noqa: F401
    import xgen_agent_runtime.llm_client.google  # noqa: F401
    import xgen_agent_runtime.llm_client.openai  # noqa: F401
    import xgen_agent_runtime.memory.factory  # noqa: F401
    import mcp.server  # noqa: F401
    import xgen_doc2chunk  # noqa: F401

    print(f"imports ok — runtime {metadata.version('xgen-agent-runtime')}, python {sys.version.split()[0]}")


def _pdf_bytes(text: str) -> bytes:
    """글자 한 줄짜리 PDF — 손으로 만든다(외부 파일 없이)."""
    stream = f"BT /F1 24 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def make_documents(folder: Path) -> dict[str, str]:
    """형식마다 표식 글자를 담은 문서를 만든다 → {파일 이름: 표식}."""
    import docx
    import openpyxl
    import pptx

    folder.mkdir(parents=True, exist_ok=True)
    marks = {}
    d = docx.Document()
    d.add_paragraph("XD-DOCX-OK")
    d.save(folder / "sample.docx")
    marks["sample.docx"] = "XD-DOCX-OK"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "XD-XLSX-OK"
    wb.save(folder / "sample.xlsx")
    marks["sample.xlsx"] = "XD-XLSX-OK"
    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "XD-PPTX-OK"
    deck.save(folder / "sample.pptx")
    marks["sample.pptx"] = "XD-PPTX-OK"
    (folder / "sample.pdf").write_bytes(_pdf_bytes("XD-PDF-OK"))
    marks["sample.pdf"] = "XD-PDF-OK"
    return marks


def check_daemon() -> None:
    tmp = Path(tempfile.mkdtemp(prefix="xd-verify-"))
    root = tmp / "XD"
    workspace = root / "workspace" / "Verify"
    marks = make_documents(workspace)
    calls = [{"name": "Write", "input": {"file_path": "notes.txt", "content": "bundle ok"}},
             {"name": "Bash", "input": {"command": "echo engine-ok"}}]
    calls += [{"name": "ParseDocument", "input": {"file_path": name}} for name in marks]
    script = tmp / "script.json"
    script.write_text(json.dumps({"responses": [{"tools": calls}, {"text": "verified"}]}), encoding="utf-8")

    env = {k: v for k, v in os.environ.items() if not k.startswith("PYTHON")}
    env["XD_ENGINE_FAKE_LLM"] = str(script)
    proc = subprocess.Popen(
        [sys.executable, "-I", "-m", "xd_engine", "--root", str(root)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=open(tmp / "engine-stderr.log", "wb"),
        env=env,
    )
    events: "queue.Queue[dict]" = queue.Queue()

    def read() -> None:
        assert proc.stdout is not None
        for raw in iter(proc.stdout.readline, b""):
            events.put(json.loads(raw.decode("utf-8")))  # 프로토콜 밖의 줄이면 여기서 깨진다

    threading.Thread(target=read, daemon=True).start()

    def send(cmd: dict) -> None:
        assert proc.stdin is not None
        proc.stdin.write((json.dumps(cmd) + "\n").encode("utf-8"))
        proc.stdin.flush()

    seen: list[dict] = []

    def until(pred, timeout: float = 120.0) -> dict:
        deadline = time.monotonic() + timeout
        while time.monotonic() < deadline:
            try:
                event = events.get(timeout=1)
            except queue.Empty:
                if proc.poll() is not None:
                    break
                continue
            seen.append(event)
            if pred(event):
                return event
        log = (tmp / "engine-stderr.log").read_text(errors="replace")[-4000:]
        raise SystemExit(f"daemon check failed; events: {seen}\nstderr:\n{log}")

    started = time.monotonic()
    ready = until(lambda e: e["type"] == "ready", timeout=60)
    print(f"ready in {time.monotonic() - started:.1f}s — protocol {ready['protocol']}, runtime {ready['runtime']}")
    send({"type": "configure", "dangerous": []})
    until(lambda e: e["type"] == "configured")
    send({
        "type": "turn", "id": "v1", "conversation": "verify", "text": "verify the bundle",
        "agent": {"id": "verify", "name": "Verify", "workspace": "Verify"},
        "config": {"provider": "xd_fake", "model": "fake", "memory_distill": False},
    })
    end = until(lambda e: e.get("id") == "v1" and e["type"] in TERMINALS)
    send({"type": "shutdown"})
    code = proc.wait(timeout=60)

    results = [e["event"] for e in seen if e["type"] == "tool" and e["event"]["type"] != "tool_call"]
    problems = []
    if end["type"] != "done":
        problems.append(f"turn ended with {end}")
    if code != 0:
        problems.append(f"daemon exit code {code}")
    if len([e for e in seen if e.get("id") == "v1" and e["type"] in TERMINALS]) != 1:
        problems.append("not exactly one terminal event")
    errors = [r for r in results if r["type"] == "tool_error"]
    if errors:
        problems.append(f"tool errors: {errors}")
    text = json.dumps(results, ensure_ascii=False)
    for name, mark in marks.items():
        if mark not in text:
            problems.append(f"ParseDocument did not read {name}")
    if "engine-ok" not in text:
        problems.append("Bash did not run")
    if (workspace / "notes.txt").read_text(encoding="utf-8") != "bundle ok":
        problems.append("Write did not land in the workspace")
    if problems:
        log = (tmp / "engine-stderr.log").read_text(errors="replace")[-4000:]
        raise SystemExit("daemon check failed:\n- " + "\n- ".join(problems) + f"\nstderr:\n{log}")
    print(f"turn ok — {len(results)} tool results, {', '.join(marks)} parsed, exit {code}")


if __name__ == "__main__":
    check_forbidden()
    check_imports()
    check_daemon()
    print("bundle verified")
